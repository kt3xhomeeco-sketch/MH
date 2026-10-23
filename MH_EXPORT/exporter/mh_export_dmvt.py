# -*- coding: utf-8 -*-
"""CLI exporter for MH_EXPORT SketchUp data.

Input JSON schema:
{
  "template_path": "D:/Cong viec/Sample/DMVT-VAN.xlsx",
  "output_path": "...xlsx",
  "proposer": "...",
  "content": "model/file name",
  "rows": [
    {"name": "Ván", "size": "17.6mm", "material_code": "...", "unit": "tấm", "quantity": 9, "note": ""}
  ]
}
"""

from __future__ import annotations

import argparse
import json
import importlib.util
import os
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side


DEFAULT_TEMPLATE = r"D:\Cong viec\Sample\DMVT-VAN.xlsx"
APP_DIR = os.path.dirname(os.path.abspath(__file__))


def load_payload(path: str) -> dict:
    with open(path, "r", encoding="utf-8-sig") as handle:
        return json.load(handle)


def set_cell_value(ws, address: str, value) -> None:
    cell = ws[address]
    if type(cell).__name__ != "MergedCell":
        cell.value = value
        return

    for merged_range in ws.merged_cells.ranges:
        if address in merged_range:
            ws.cell(row=merged_range.min_row, column=merged_range.min_col).value = value
            return


def border_range(ws, min_row: int, max_row: int, min_col: int = 1, max_col: int = 8) -> None:
    thin = Side(border_style="thin", color="000000")
    border_style = Border(top=thin, left=thin, right=thin, bottom=thin)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = border_style


def plain_range(ws, min_row: int, max_row: int, min_col: int = 1, max_col: int = 8) -> None:
    no_fill = PatternFill(fill_type=None)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.font = Font(name=cell.font.name or "Arial", size=cell.font.sz or 11, bold=False)
            cell.fill = no_fill


def clear_data_rows(ws, start_row: int = 11) -> None:
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)


def normalize_quantity(value):
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip().replace(",", ".")
    try:
        number = float(text)
        return int(number) if number.is_integer() else number
    except Exception:
        return value


def display_row_name(row: dict) -> str:
    name = str(row.get("name", "")).strip()
    return name or "Ván"


def write_rows(ws, rows: list[dict], start_row: int = 11) -> int | None:
    if not rows:
        return None

    directory = Path(os.environ.get("MH_NHAP_XUATVAN_LIBRARY", r"D:\Cong viec\config\nhap_xuatvan"))
    spec = importlib.util.spec_from_file_location("mh_export_size_rules", directory / "material_item_names.py")
    size_rules = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(size_rules)
    for offset, row in enumerate(rows):
        excel_row = start_row + offset
        size = row.get("size", "")
        if str(row.get("unit", "")).strip().lower() in ("tấm", "tam", "m2", "m²"):
            size = size_rules.material_code_size(row.get("material_code", ""), fallback=size)
        values = [
            offset + 1,
            display_row_name(row),
            size,
            row.get("material_code", ""),
            row.get("unit", ""),
            normalize_quantity(row.get("quantity", "")),
            row.get("order", ""),
            row.get("note", ""),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=excel_row, column=col, value=value)

    last_row = start_row + len(rows) - 1
    plain_range(ws, start_row, last_row, 1, 8)
    border_range(ws, start_row, last_row, 1, 8)
    return last_row


def recreate_sheet(wb, title: str):
    if title in wb.sheetnames:
        del wb[title]
    return wb.create_sheet(title)


def write_abf_data_sheet(wb, abf_data: dict | None) -> None:
    if not abf_data:
        return

    headers = list(abf_data.get("headers") or [])
    rows = list(abf_data.get("rows") or [])
    if not headers and not rows:
        return

    ws = recreate_sheet(wb, "ABF Data")
    for col, value in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=value)
    for row_index, row in enumerate(rows, start=2):
        for col, value in enumerate(list(row), start=1):
            ws.cell(row=row_index, column=col, value=value)

    max_col = max(len(headers), max((len(list(row)) for row in rows), default=0), 1)
    plain_range(ws, 1, max(len(rows) + 1, 1), 1, max_col)
    border_range(ws, 1, max(len(rows) + 1, 1), 1, max_col)


def build_workbook(payload: dict, template_path: str):
    if template_path and os.path.exists(template_path):
        wb = load_workbook(template_path)
        ws = wb.active
        ws.title = "Xuất ván"
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Xuất ván"
        headers = ["STT", "Tên hàng", "Kích thước", "Mã hiệu", "Đơn vị", "Số lượng", "Hợp đồng/ Đơn hàng", "Ghi chú"]
        for col, header in enumerate(headers, start=1):
            ws.cell(row=10, column=col, value=header)

    clear_data_rows(ws, 11)
    plain_range(ws, 10, 10, 1, 8)
    border_range(ws, 10, 10, 1, 8)

    proposer = payload.get("proposer", "")
    content = payload.get("content", "")
    today = datetime.now().strftime("%d/%m/%Y")
    set_cell_value(ws, "C6", proposer)
    set_cell_value(ws, "C8", today)
    set_cell_value(ws, "C9", content)

    write_rows(ws, list(payload.get("rows") or []), 11)
    write_abf_data_sheet(wb, payload.get("abf_excel_data"))
    return wb


def export_dmvt(input_json: str, output_path: str, template_path: str | None) -> dict:
    payload = load_payload(input_json)
    template = template_path or payload.get("template_path") or DEFAULT_TEMPLATE
    output = output_path or payload.get("output_path")
    if not output:
        raise ValueError("Missing output path")

    os.makedirs(os.path.dirname(os.path.abspath(output)), exist_ok=True)
    wb = build_workbook(payload, template)
    wb.save(output)
    return {
        "ok": True,
        "output_path": output,
        "row_count": len(payload.get("rows") or []),
        "message": "Exported DMVT workbook",
    }


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export MH DMVT Excel from SketchUp JSON")
    parser.add_argument("--input", required=True, help="Input JSON path")
    parser.add_argument("--output", required=True, help="Output .xlsx path")
    parser.add_argument("--template", default=DEFAULT_TEMPLATE, help="DMVT template .xlsx path")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    try:
        result = export_dmvt(args.input, args.output, args.template)
        print(json.dumps(result, ensure_ascii=False))
        return 0
    except Exception as exc:
        print(json.dumps({"ok": False, "message": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
