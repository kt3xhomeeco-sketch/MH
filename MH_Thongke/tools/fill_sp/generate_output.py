from openpyxl import load_workbook
from copy import copy
import argparse
import os
import sys
from pathlib import Path
import threading
import traceback
import unicodedata
import re

HEADER_ROW = 7
DATA_START_ROW = 9


def norm(v):
    if v is None:
        return ''
    return str(v).strip()


def normalize_text(v):
    s = norm(v).lower()
    s = unicodedata.normalize('NFD', s)
    s = ''.join(ch for ch in s if unicodedata.category(ch) != 'Mn')
    s = s.replace('đ', 'd')
    return ' '.join(s.split())


def is_section_header(ma_sp_raw, ten_sp_raw, c5, c6, c7):
    if c5 or c6 or c7 or not ten_sp_raw:
        return False
    if ma_sp_raw and len(ma_sp_raw) == 1 and ma_sp_raw.isalpha():
        return True
    upper_name = ten_sp_raw.upper()
    return upper_name.startswith('PHÒNG') or upper_name.startswith('WC') or upper_name == 'SẢNH'


def find_header_row(ws, required_columns, max_rows=24):
    for row_num in range(1, min(ws.max_row, max_rows) + 1):
        header = {normalize_text(ws.cell(row_num, c).value): c for c in range(1, ws.max_column + 1)}
        resolved = {}
        for key, aliases in required_columns.items():
            resolved[key] = next((header.get(alias) for alias in aliases if header.get(alias)), None)
        if all(resolved.values()):
            return row_num, header, resolved
    return None, {}, {}


def normalized_group(value):
    return re.sub(r'^\s*[1-4]\s+', '', normalize_text(value)).strip()


def is_disabled(value):
    return normalize_text(value) in ('0', 'false', 'no', 'khong', 'inactive', 'off')


def source_columns(value, fallback=('F', 'G')):
    columns = re.findall(r'\b[A-Z]{1,3}\b', str(value or '').upper())
    result = []
    for column in columns:
        number = 0
        for letter in column:
            number = (number * 26) + (ord(letter) - ord('A') + 1)
        if 1 <= number <= 16384 and column not in result:
            result.append(column)
    return result or list(fallback)


def source_column_number(column):
    number = 0
    for letter in str(column or '').upper():
        number = (number * 26) + (ord(letter) - ord('A') + 1)
    return number


def library_generation_rules(wb):
    rules = {'segment_separator': '.', 'multi_separator': '&', 'source_columns': ['F', 'G']}
    if 'material_code_rules' not in wb.sheetnames:
        return rules

    ws = wb['material_code_rules']
    row_num, _header, columns = find_header_row(ws, {
        'name': ('quy tac', 'rule'),
        'value': ('gia tri', 'value')
    })
    if not row_num:
        return rules

    for r in range(row_num + 1, ws.max_row + 1):
        name = normalize_text(ws.cell(r, columns['name']).value)
        value = norm(ws.cell(r, columns['value']).value)
        if name == 'dau ghep cac phan' and value:
            rules['segment_separator'] = value
        elif name == 'dau ngan nhieu vat lieu' and value:
            rules['multi_separator'] = value
        elif name in ('cot excel doc de tao ma', 'cot du lieu doc de tao ma', 'cot nguon tao ma'):
            rules['source_columns'] = source_columns(value)
    return rules


def load_mapping(path):
    wb = load_workbook(path, data_only=True)
    ws = wb['mapping'] if 'mapping' in wb.sheetnames else wb[wb.sheetnames[0]]
    header_row, header, columns = find_header_row(ws, {
        'raw': ('tu khoa trong excel', 'gia tri', 'gia_tri', 'raw value', 'raw_value', 'input'),
        'mapped': ('ma tao ra', 'gia tri map', 'gia_tri_map', 'mapping', 'ma map', 'ma_map')
    })
    if not header_row:
        raise ValueError('Sheet mapping phải có cột `Từ khóa trong Excel` và `Mã tạo ra`.')

    raw_col = columns['raw']
    map_col = columns['mapped']
    group_col = header.get('nhom') or header.get('group')
    enabled_col = header.get('dang dung') or header.get('enabled') or header.get('active')
    order_col = header.get('thu tu') or header.get('thu_tu') or header.get('order')

    rules = []
    for r in range(header_row + 1, ws.max_row + 1):
        raw_text = norm(ws.cell(r, raw_col).value)
        mapped = norm(ws.cell(r, map_col).value)
        if not raw_text or not mapped:
            continue
        if enabled_col and is_disabled(ws.cell(r, enabled_col).value):
            continue

        group_name = normalized_group(ws.cell(r, group_col).value) if group_col else ''
        order_raw = ws.cell(r, order_col).value if order_col else None
        try:
            order_val = int(order_raw) if order_raw is not None and str(order_raw).strip() != '' else None
        except (ValueError, TypeError):
            order_val = None

        rules.append({
            'raw': raw_text,
            'raw_norm': normalize_text(raw_text),
            'mapped': mapped,
            'group': group_name,
            'order': order_val,
        })
        if normalize_text(raw_text) == 'mdf chong am':
            rules.append({
                'raw': 'MDF CA',
                'raw_norm': 'mdf ca',
                'mapped': mapped,
                'group': group_name,
                'order': order_val,
            })

    rules.sort(key=lambda x: len(x['raw_norm']), reverse=True)
    return rules, library_generation_rules(wb)


def group_rank(g):
    g = normalized_group(g)
    if g in ('chat lieu', 'chat_lieu', 'material'):
        return 0
    if g in ('nha cung cap', 'nha_cung_cap', 'supplier'):
        return 1
    if g in ('be mat', 'be_mat', 'surface'):
        return 2
    if g in ('ma vat lieu', 'ma_vat_lieu', 'material code', 'finish code', 'code'):
        return 3
    return 4


def remove_redundant_material_codes(codes):
    # Bỏ mã ngắn nếu đã có mã dài cùng prefix: HDFCA giữ, bỏ HDF
    result = []
    for c in codes:
        redundant = False
        for other in codes:
            if c != other and other.startswith(c) and len(other) > len(c):
                redundant = True
                break
        if not redundant:
            result.append(c)
    return result


def remove_redundant_surface_codes(codes):
    result = []
    for code in codes:
        redundant = False
        for other in codes:
            if code != other and other.startswith(code) and len(other) > len(code):
                redundant = True
                break
        if not redundant:
            result.append(code)
    return result


def span_covered(span, covered_spans):
    start, end = span
    return any(start >= a and end <= b for a, b in covered_spans)


def extract_dynamic_surface_codes(src_norm):
    """Đọc mã LK/MS nằm trong mô tả và trả về một token hoàn chỉnh.

    Ví dụ: ``Laminate LK 1201 T`` -> ``LK1201T``.  Không tách thành
    ``LK.1201T`` vì Desktop và Material Assign dùng một mã hoàn thiện liền.
    """
    codes = []
    for match in re.finditer(r'\b(MS|LK)\s*([A-Z0-9]+)(?:\s+([A-Z0-9]+))?', src_norm, flags=re.IGNORECASE):
        first = match.group(2).upper()
        second = match.group(3).upper() if match.group(3) else ''
        if not re.search(r'\d', first) and not re.search(r'\d', second):
            continue
        code = f"{match.group(1).upper()}{first}{second}"
        if code not in codes:
            codes.append(code)

    if codes:
        return codes

    # Mô tả cũ không ghi rõ prefix LK/MS.
    m = re.search(r'laminate\s+ma\s+([a-z]{1,3}\s*)?([a-z0-9]{3,10})', src_norm, flags=re.IGNORECASE)
    if m:
        codes.append(f"LK{m.group(2).upper()}")
    m = re.search(r'melamin(?:e)?.*?\sma\s+([a-z]{1,3}\s*)?([a-z0-9]{2,10})(?:\s+([a-z]{1,4}))?', src_norm, flags=re.IGNORECASE)
    if m:
        raw = m.group(2).upper()
        if m.group(3) and raw.isdigit():
            raw = f"{raw}{m.group(3).upper()}"
        if raw.startswith('MS') and len(raw) > 2:
            raw = raw[2:]
        codes.append(f"MS{raw}")
    return list(dict.fromkeys(codes))


def extract_dynamic_veneer_codes(src_norm):
    # PHỦ VENEER VHT 52 AN CƯỜNG => VER.VHT52
    m = re.search(r'veneer\s+([a-z]{2,5})\s*([0-9]{1,4})\b', src_norm, flags=re.IGNORECASE)
    if not m:
        return []
    return ['VER', f"{m.group(1)}{m.group(2)}".upper()]


def extract_son_code(ht_text):
    # MÀU DULUX 20YY 74/055 => 20YY-74/055
    # MÀU DULUX 20YY 74/055 20% => 20YY-74/055-20
    raw = norm(ht_text).upper()
    raw_norm = normalize_text(raw)

    m = re.search(r'DULUX\s+([0-9A-Z]+)\s+([0-9]{2,3}/[0-9]{2,3})(?:.*?([0-9]{1,3})\s*%)?', raw)
    if m:
        code = f"{m.group(1)}-{m.group(2)}"
        if m.group(3):
            code = f"{code}-{m.group(3)}"
        return code

    m = re.search(r'\b([0-9]{2}[A-Z]{2})\s+([0-9]{2,3}/[0-9]{2,3})(?:.*?([0-9]{1,3})\s*%)?', raw)
    if m:
        code = f"{m.group(1)}-{m.group(2)}"
        if m.group(3):
            code = f"{code}-{m.group(3)}"
        return code

    m = re.search(r'\b(BQ[0-9]+(?:/[0-9]+)?)(?:.*?([0-9]{1,3})\s*%)?', raw)
    if m:
        code = m.group(1)
        if m.group(2):
            code = f"{code}-{m.group(2)}"
        return code

    if raw_norm == 'son':
        return 'SON'
    if raw_norm == 'son bet':
        return 'SONBET'

    return ''


def build_ma_vat_lieu(qc_text, ht_text, rules, include_son_code=False, segment_separator='.', multi_separator='&'):
    src_norm = normalize_text(f"{norm(qc_text)} {norm(ht_text)}")
    if not src_norm:
        return ''

    found = []
    used_codes = set()
    covered_spans = []
    for rule in rules:
        token = rule['raw_norm']
        pos = src_norm.find(token) if token else -1
        if pos < 0:
            continue

        span = (pos, pos + len(token))
        if span_covered(span, covered_spans):
            continue

        if rule['mapped'] not in used_codes:
            item = rule.copy()
            item['pos'] = pos
            found.append(item)
            used_codes.add(rule['mapped'])
            covered_spans.append(span)

    if not found:
        found = []

    found.sort(key=lambda x: (x['order'] if x['order'] is not None else 999, group_rank(x['group']), -len(x['raw_norm'])))

    material_rules = sorted(
        [x for x in found if group_rank(x['group']) == 0 and x['mapped']],
        key=lambda x: (x['pos'] if x.get('pos', -1) >= 0 else 999999, -len(x['raw_norm']))
    )
    material_codes = [x['mapped'] for x in material_rules]
    supplier_codes = [x['mapped'] for x in found if group_rank(x['group']) == 1 and x['mapped']]
    surface_codes = [x['mapped'] for x in found if group_rank(x['group']) == 2 and x['mapped']]
    finish_codes = [x['mapped'] for x in found if group_rank(x['group']) == 3 and x['mapped']]

    material_codes = remove_redundant_material_codes(material_codes)
    material_codes = list(dict.fromkeys(material_codes))
    supplier_codes = list(dict.fromkeys(supplier_codes))
    surface_codes = remove_redundant_surface_codes(list(dict.fromkeys(surface_codes)))
    finish_codes = remove_redundant_surface_codes(list(dict.fromkeys(finish_codes)))

    # Mã động cho laminate / melamine (không cần khai báo trong library)
    for dyn_code in extract_dynamic_surface_codes(src_norm):
        prefix = dyn_code[:2]
        # Khi có mã chi tiết LK/MS, giữ LAM/MEL khai báo trong thư viện nhưng
        # thay mọi mã động cùng prefix bằng token hoàn chỉnh mới nhất.
        finish_codes = [c for c in finish_codes if not c.upper().startswith(prefix)]
        finish_codes.append(dyn_code)

    veneer_codes = extract_dynamic_veneer_codes(src_norm)
    for code in veneer_codes:
        if code not in surface_codes and code not in finish_codes:
            finish_codes.append(code)
    surface_codes = remove_redundant_surface_codes(list(dict.fromkeys(surface_codes)))
    finish_codes = remove_redundant_surface_codes(list(dict.fromkeys(finish_codes)))

    # TẦN BÌ KẾT HỢP PLYWOOD => TANBI & PLY
    if len(material_codes) >= 2:
        left = f' {multi_separator} '.join(material_codes)
    elif len(material_codes) == 1:
        left = material_codes[0]
    else:
        left = ''

    right_parts = supplier_codes + surface_codes + finish_codes
    son_code = extract_son_code(ht_text or qc_text) if include_son_code else ''
    # "Sơn" có thể đã được mapping thành SON. Không nối thêm lần nữa.
    if son_code and normalize_text(son_code) not in [normalize_text(code) for code in right_parts]:
        right_parts.append(son_code)

    if right_parts:
        right = segment_separator.join(right_parts)
        return f"{left}{segment_separator}{right}" if left else right
    return left


def fill_file(input_xlsx, output_xlsx, lib_xlsx, include_son_code=True):
    rules, library_rules = load_mapping(lib_xlsx)
    material_source_columns = source_columns(library_rules.get('source_columns'))

    wb = load_workbook(input_xlsx)
    ws = wb[wb.sheetnames[0]]

    src = ws.cell(HEADER_ROW, 8)
    for c, t in [(9, 'Ma_SP'), (10, 'Ten_SP'), (11, 'Ma_vat_lieu')]:
        cell = ws.cell(HEADER_ROW, c, t)
        if src.has_style:
            cell._style = copy(src._style)
            cell.font = copy(src.font)
            cell.fill = copy(src.fill)
            cell.border = copy(src.border)
            cell.alignment = copy(src.alignment)
            cell.number_format = src.number_format
            cell.protection = copy(src.protection)

    current_ma_sp = ''
    current_ten_sp = ''

    for r in range(DATA_START_ROW, ws.max_row + 1):
        ma_sp_raw = norm(ws.cell(r, 1).value)
        ten_sp_raw = norm(ws.cell(r, 2).value)
        c5 = norm(ws.cell(r, 5).value)
        c6 = norm(ws.cell(r, 6).value)
        c7 = norm(ws.cell(r, 7).value)

        if ma_sp_raw.upper().startswith('SP'):
            current_ma_sp = ma_sp_raw
            current_ten_sp = ten_sp_raw
        elif ten_sp_raw:
            if is_section_header(ma_sp_raw, ten_sp_raw, c5, c6, c7):
                current_ma_sp = ''
                current_ten_sp = ''
            else:
                current_ten_sp = ten_sp_raw

        ws.cell(r, 9, current_ma_sp or None)
        ws.cell(r, 10, current_ten_sp or None)

        material_source = ' '.join(
            norm(ws.cell(r, source_column_number(column)).value)
            for column in material_source_columns
            if norm(ws.cell(r, source_column_number(column)).value)
        )
        mvl = build_ma_vat_lieu(
            material_source,
            '',
            rules,
            include_son_code=include_son_code,
            segment_separator=library_rules['segment_separator'],
            multi_separator=library_rules['multi_separator']
        )
        ws.cell(r, 11, mvl or None)

    # Keep every source row intact.  Generated codes are filled per row; rows
    # with the same column F/value must not be merged or deleted.
    wb.save(output_xlsx)


def choose_input_file():
    try:
        from tkinter import Tk, filedialog

        root = Tk()
        root.withdraw()
        root.attributes('-topmost', True)
        path = filedialog.askopenfilename(
            title='Chọn file Excel cần xử lý',
            filetypes=[('Excel files', '*.xlsx *.xlsm'), ('All files', '*.*')],
        )
        root.destroy()
        return path
    except Exception:
        return input('Nhập hoặc kéo thả đường dẫn file Excel vào đây: ').strip().strip('"')


def default_output_path(input_xlsx, include_son_code=True):
    src = Path(input_xlsx)
    suffix = '_generated_son' if include_son_code else '_generated'
    return str(src.with_name(f'{src.stem}{suffix}{src.suffix}'))


def run_gui():
    from tkinter import BooleanVar, StringVar, Tk, filedialog, messagebox
    from tkinter import ttk

    root = Tk()
    root.title('Generate Ma_SP / Ten_SP / Ma_vat_lieu')
    root.geometry('760x360')
    root.minsize(720, 330)

    input_var = StringVar()
    output_var = StringVar()
    library_var = StringVar(value=str(Path('library.xlsx').resolve()))
    include_son_var = BooleanVar(value=True)
    status_var = StringVar(value='Chọn file Excel rồi bấm Chạy.')

    def refresh_output_path(*_):
        if input_var.get().strip() and not output_var.get().strip():
            output_var.set(default_output_path(input_var.get().strip(), include_son_var.get()))

    def browse_input():
        path = filedialog.askopenfilename(
            title='Chọn file Excel cần xử lý',
            filetypes=[('Excel files', '*.xlsx *.xlsm'), ('All files', '*.*')],
        )
        if path:
            input_var.set(path)
            output_var.set(default_output_path(path, include_son_var.get()))

    def browse_output():
        initial = output_var.get().strip() or default_output_path(input_var.get().strip(), include_son_var.get())
        path = filedialog.asksaveasfilename(
            title='Chọn nơi lưu file output',
            initialfile=Path(initial).name if initial else 'output_generated_son.xlsx',
            defaultextension='.xlsx',
            filetypes=[('Excel files', '*.xlsx'), ('All files', '*.*')],
        )
        if path:
            output_var.set(path)

    def browse_library():
        path = filedialog.askopenfilename(
            title='Chọn library.xlsx',
            initialfile='library.xlsx',
            filetypes=[('Excel files', '*.xlsx *.xlsm'), ('All files', '*.*')],
        )
        if path:
            library_var.set(path)

    def open_output():
        path = output_var.get().strip()
        if not path or not Path(path).exists():
            messagebox.showwarning('Chưa có file', 'Chưa có file output để mở.')
            return
        os.startfile(path)

    def set_busy(is_busy):
        state = 'disabled' if is_busy else 'normal'
        run_button.config(state=state)
        open_button.config(state='disabled' if is_busy else 'normal')

    def append_log(text):
        log_box.config(state='normal')
        log_box.insert('end', text + '\n')
        log_box.see('end')
        log_box.config(state='disabled')

    def run_job():
        input_path = input_var.get().strip()
        output_path = output_var.get().strip()
        library_path = library_var.get().strip()

        if not input_path:
            messagebox.showwarning('Thiếu input', 'Chọn file Excel đầu vào trước.')
            return
        if not Path(input_path).exists():
            messagebox.showerror('Không tìm thấy input', input_path)
            return
        if not library_path:
            messagebox.showwarning('Thiếu library', 'Chọn file library.xlsx trước.')
            return
        if not Path(library_path).exists():
            messagebox.showerror('Không tìm thấy library', library_path)
            return
        if not output_path:
            output_path = default_output_path(input_path, include_son_var.get())
            output_var.set(output_path)

        def worker():
            root.after(0, set_busy, True)
            root.after(0, status_var.set, 'Đang xử lý...')
            root.after(0, append_log, f'Input: {input_path}')
            root.after(0, append_log, f'Output: {output_path}')
            try:
                fill_file(input_path, output_path, library_path, include_son_code=include_son_var.get())
            except Exception:
                err = traceback.format_exc()
                root.after(0, append_log, err)
                root.after(0, status_var.set, 'Lỗi. Xem log bên dưới.')
                root.after(0, messagebox.showerror, 'Lỗi xử lý', err)
            else:
                root.after(0, append_log, 'Done.')
                root.after(0, status_var.set, f'Đã tạo: {output_path}')
                root.after(0, messagebox.showinfo, 'Xong', f'Đã tạo file:\n{output_path}')
            finally:
                root.after(0, set_busy, False)

        threading.Thread(target=worker, daemon=True).start()

    root.columnconfigure(1, weight=1)
    root.rowconfigure(6, weight=1)

    pad = {'padx': 10, 'pady': 6}
    ttk.Label(root, text='File input').grid(row=0, column=0, sticky='w', **pad)
    ttk.Entry(root, textvariable=input_var).grid(row=0, column=1, sticky='ew', **pad)
    ttk.Button(root, text='Chọn...', command=browse_input).grid(row=0, column=2, sticky='ew', **pad)

    ttk.Label(root, text='File output').grid(row=1, column=0, sticky='w', **pad)
    ttk.Entry(root, textvariable=output_var).grid(row=1, column=1, sticky='ew', **pad)
    ttk.Button(root, text='Lưu thành...', command=browse_output).grid(row=1, column=2, sticky='ew', **pad)

    ttk.Label(root, text='Library').grid(row=2, column=0, sticky='w', **pad)
    ttk.Entry(root, textvariable=library_var).grid(row=2, column=1, sticky='ew', **pad)
    ttk.Button(root, text='Chọn...', command=browse_library).grid(row=2, column=2, sticky='ew', **pad)

    ttk.Checkbutton(
        root,
        text='Xuất cả mã sơn',
        variable=include_son_var,
        command=lambda: output_var.set(default_output_path(input_var.get().strip(), include_son_var.get())) if input_var.get().strip() else None,
    ).grid(row=3, column=1, sticky='w', **pad)

    action_frame = ttk.Frame(root)
    action_frame.grid(row=4, column=0, columnspan=3, sticky='w', **pad)
    run_button = ttk.Button(action_frame, text='Chạy', command=run_job)
    run_button.grid(row=0, column=0, padx=(0, 8))
    open_button = ttk.Button(action_frame, text='Mở output', command=open_output)
    open_button.grid(row=0, column=1)

    ttk.Label(root, textvariable=status_var).grid(row=5, column=0, columnspan=3, sticky='w', padx=10, pady=6)

    from tkinter import Text
    log_box = Text(root, height=8, wrap='word', state='disabled')
    log_box.grid(row=6, column=0, columnspan=3, sticky='nsew', padx=10, pady=(6, 10))

    input_var.trace_add('write', refresh_output_path)
    root.mainloop()


def main():
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')

    parser = argparse.ArgumentParser(description='Generate output excel with Ma_SP, Ten_SP, Ma_vat_lieu')
    parser.add_argument('--input')
    parser.add_argument('--output')
    parser.add_argument('--library', default='library.xlsx')
    parser.add_argument('--include-son-code', action='store_true', default=True, help='Append Dulux paint code into Ma_vat_lieu')
    parser.add_argument('--no-son-code', dest='include_son_code', action='store_false', help='Do not append Dulux paint code into Ma_vat_lieu')
    args = parser.parse_args()

    if len(sys.argv) == 1:
        run_gui()
        return

    if not args.input:
        args.input = choose_input_file()
    if not args.input:
        print('Canceled: chưa chọn file input.')
        sys.exit(1)
    if not args.output:
        args.output = default_output_path(args.input, include_son_code=args.include_son_code)

    fill_file(args.input, args.output, args.library, include_son_code=args.include_son_code)
    print(f'Created: {args.output}')


if __name__ == '__main__':
    main()
